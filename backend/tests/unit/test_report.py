from __future__ import annotations

import csv

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from dxa_qc.domain import AnatomicalRegion
from dxa_qc.report import REPORT_COLUMNS, ProcessingStatus, ReportFormat, ResultRow, write_report


def ok_row(**kw) -> ResultRow:
    base = dict(
        path_to_study="batch/s1/a.dcm",
        study_uid="1.2.3",
        image_uid="1.2.3.4",
        anatomical_region=AnatomicalRegion.SPINE,
        quality_class=1,
        quality_prob=0.87654321,
        violation_type="Некорректная укладка",
        processing_status=ProcessingStatus.SUCCESS,
        time_of_processing=1.23456,
    )
    return ResultRow(**{**base, **kw})


def fail_row(**kw) -> ResultRow:
    base = dict(
        path_to_study="batch/s1/broken.dcm",
        processing_status=ProcessingStatus.FAILURE,
        time_of_processing=0.01,
        error="not a valid DICOM file",
    )
    return ResultRow(**{**base, **kw})


# --- schema invariants ---


def test_success_row_requires_prediction():
    with pytest.raises(ValidationError, match="needs anatomical_region"):
        ok_row(quality_prob=None)


def test_class_zero_with_violation_is_contradiction():
    with pytest.raises(ValidationError, match="contradicts"):
        ok_row(quality_class=0)


def test_failure_requires_error():
    with pytest.raises(ValidationError, match="explain"):
        fail_row(error=None)


@pytest.mark.parametrize(
    "field,value", [("quality_prob", 1.01), ("quality_prob", -0.1), ("quality_class", 2)]
)
def test_ranges(field, value):
    with pytest.raises(ValidationError):
        ok_row(**{field: value})


# --- writers ---


def test_xlsx_columns_types_and_values(tmp_path):
    out = write_report([ok_row(), fail_row()], tmp_path, ReportFormat.XLSX)
    ws = load_workbook(out.table).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == REPORT_COLUMNS
    good = dict(zip(REPORT_COLUMNS, rows[1], strict=True))
    bad = dict(zip(REPORT_COLUMNS, rows[2], strict=True))
    assert good["anatomical_region"] == "Поясничный отдел позвоночника"
    assert good["quality_class"] == 1 and isinstance(good["quality_class"], int)
    assert good["quality_prob"] == pytest.approx(0.876543)
    assert good["time_of_processing"] == pytest.approx(1.235)
    assert bad["processing_status"] == "Failure"
    assert bad["quality_class"] is None and bad["anatomical_region"] is None


def test_csv_roundtrip_utf8(tmp_path):
    out = write_report([ok_row(), fail_row()], tmp_path, ReportFormat.CSV)
    with out.table.open(encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert list(rows[0].keys()) == list(REPORT_COLUMNS)
    assert rows[0]["violation_type"] == "Некорректная укладка"
    assert rows[1]["quality_class"] == ""


def test_errors_file_only_when_failures(tmp_path):
    assert write_report([ok_row()], tmp_path / "a", ReportFormat.CSV).errors is None
    out = write_report([ok_row(), fail_row()], tmp_path / "b", ReportFormat.CSV)
    with out.errors.open(encoding="utf-8") as fh:
        assert list(csv.reader(fh)) == [
            ["path_to_study", "error"],
            ["batch/s1/broken.dcm", "not a valid DICOM file"],
        ]


def test_error_text_not_in_main_table(tmp_path):
    out = write_report([fail_row(error="SECRET_DETAIL")], tmp_path, ReportFormat.CSV)
    assert "SECRET_DETAIL" not in out.table.read_text(encoding="utf-8")


@pytest.mark.parametrize("fmt", list(ReportFormat))
def test_formula_injection_neutralized(tmp_path, fmt):
    evil = '=HYPERLINK("http://evil","x").dcm'
    out = write_report([fail_row(path_to_study=evil)], tmp_path, fmt)
    if fmt is ReportFormat.XLSX:
        cell = load_workbook(out.table).active.cell(row=2, column=1)
        assert cell.data_type == "s"
        assert cell.value == "'" + evil
    else:
        with out.table.open(encoding="utf-8", newline="") as fh:
            assert next(csv.DictReader(fh))["path_to_study"] == "'" + evil


def test_empty_batch_writes_header_only(tmp_path):
    out = write_report([], tmp_path, ReportFormat.XLSX)
    assert list(load_workbook(out.table).active.iter_rows(values_only=True)) == [REPORT_COLUMNS]
